"""Parse fetched public data. Preserve source date, market type and units."""
import json,re,pathlib,datetime,hashlib
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from collect import ROOT,URLS
TODAY=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5))).date().isoformat()
def slug(s):return re.sub('[^a-z0-9]+','-',str(s).lower()).strip('-')
def num(s):
 m=re.search(r'\d[\d,]*(?:\.\d+)?',str(s));return float(m[0].replace(',','')) if m else None
def date_en(s):
 for fmt in ['%B %d, %Y','%b %d, %Y','%b %d, %y','%d %b %y']:
  try:return datetime.datetime.strptime(s.strip(),fmt).date().isoformat()
  except ValueError:pass
 raise ValueError('Unknown source date '+s)
def soup(key):return BeautifulSoup((ROOT/'raw'/f'{key}.html').read_text(),'html.parser')
from taxonomy import survey_category as classified
def ingest():
 records={r['id']:r for r in json.loads((ROOT/'data/prices.json').read_text())}
 hist=json.loads((ROOT/'data/history.json').read_text()) if (ROOT/'data/history.json').exists() else {}
 statuses=[]
 def add(id,name,category,price,unit,date,source,city='Pakistan',kind='Retailer listing',**kw):
  if not isinstance(price,(int,float)) or not 0<price<1e10 or date>TODAY:return
  r=dict(id=id,name=name,category=category,price=round(price,4),unit=unit,date=date,source=source,city=city,kind=kind,**kw)
  old=records.get(id)
  if not old or date>=old['date']:records[id]=r
  pts={x['date']:x for x in hist.get(id,[])};pts[date]={'date':date,'price':r['price'],'source':source};hist[id]=sorted(pts.values(),key=lambda x:x['date'])[-400:]
 def attempt(name,fn):
  try:before=len(records);fn();statuses.append({'source':name,'status':'parsed','new_records':len(records)-before})
  except Exception as e:statuses.append({'source':name,'status':'manual_required','reason':str(e)})
 def fuel():
  ss=soup('pso-history')
  for li in ss.select('li'):
   title=li.select_one('.uk-accordion-title');table=li.select_one('table')
   if not title or not table:continue
   d=re.search(r'Effective From:\s*([A-Za-z]+\s+\d+,\s*\d{4})',title.get_text(' ',strip=True))
   if not d:continue
   date=date_en(d[1])
   for tr in table.select('tbody tr'):
    td=tr.select('th,td')
    if len(td)!=2:continue
    name=td[0].get_text(' ',strip=True);v=num(td[1].get_text())
    id={'PREMIER EURO 5':'petrol','HI-CETANE DIESEL EURO 5':'diesel'}.get(name,'fuel-'+slug(name))
    add(id,name,'fuel',v,'litre',date,URLS['pso-history'],kind='PSO published rate')
 attempt('PSO dated fuel archive',fuel)
 def metals():
  ss=soup('gold');text=ss.get_text(' ',strip=True);m=re.search(r'Updated:\s*([A-Za-z]+\s+\d+,\s*\d{4})',text)
  if not m:raise ValueError('Gold source date missing')
  date=date_en(m[1])
  for t in ss.select('table'):
   headers=t.select_one('tr').get_text(' ',strip=True)
   for tr in t.select('tr')[1:]:
    td=tr.select('th,td')
    if len(td)<3:continue
    if headers.startswith('Purity Tola'):
     purity=td[0].get_text(strip=True)
     if not re.fullmatch(r'\d+K',purity):continue
     for j,unit in [(1,'tola'),(2,'10 grams')]:
      id='gold-'+purity.lower() if j==1 else 'gold-'+purity.lower()+'-10g'
      add(id,'Gold · '+purity,'metals',num(td[j].get_text()),unit,date,URLS['gold'],kind='Reported market benchmark')
    if headers.startswith('Date 24K'):
     try:d=date_en(td[0].get_text(' ',strip=True))
     except ValueError:continue
     for j,purity in [(1,'24k'),(2,'22k')]:add('gold-'+purity,'Gold · '+purity.upper(),'metals',num(td[j].get_text()),'tola',d,URLS['gold'],kind='Reported market benchmark')
  ss=soup('silver');m=re.search(r'Updated:\s*([A-Za-z]+\s+\d+,\s*\d{4})',ss.get_text(' ',strip=True))
  if not m:raise ValueError('Silver source date missing')
  d=date_en(m[1]);tables=ss.select('table')
  for t in tables:
   h=t.select_one('tr').get_text(' ',strip=True)
   for tr in t.select('tr')[1:]:
    td=tr.select('th,td')
    if len(td)<3:continue
    if h.startswith('City') and td[0].get_text(strip=True)=='Karachi':
     add('silver','Silver · reported benchmark','metals',num(td[1].get_text()),'tola',d,URLS['silver'],kind='Reported market benchmark')
     add('silver-10g','Silver · reported benchmark','metals',num(td[2].get_text()),'10 grams',d,URLS['silver'],kind='Reported market benchmark')
    if h.startswith('Date'):
     try:date=date_en(td[0].get_text(' ',strip=True))
     except ValueError:continue
     add('silver','Silver · reported benchmark','metals',num(td[1].get_text()),'tola',date,URLS['silver'],kind='Reported market benchmark')
 attempt('Published gold and silver benchmarks',metals)
 def forex():
  ss=soup('forex');m=re.search(r'Updated:\s*([A-Za-z]+\s+\d+,\s*\d{4})',ss.get_text(' ',strip=True))
  if not m:raise ValueError('Currency source date missing')
  d=date_en(m[1])
  for t in ss.select('table'):
   h=t.select_one('tr').get_text(' ',strip=True)
   market='Open market' if h=='Currency Buying Selling' else 'Interbank' if h.startswith('Currency Symbol') else None
   if not market:continue
   for tr in t.select('tr')[1:]:
    td=tr.select('th,td')
    if len(td)<3:continue
    name=td[0].get_text(' ',strip=True)
    if not name or 'Source:' in name:continue
    offset=2 if market=='Interbank' else 1
    if len(td)<=offset+1:continue
    for j,side in [(offset,'buy'),(offset+1,'sell')]:
     add('fx-'+slug(name)+'-'+slug(market)+'-'+side,name+' · '+market+' '+side,'currency',num(td[j].get_text()),'1 currency unit',d,URLS['forex'],kind='Reported '+market.lower()+' quote',side=side)
 attempt('Published currency tables',forex)
 def retail():
  count=0
  # Remove first-version category snapshot; canonical product URL becomes ID.
  records.pop('samsung-a06',None)
  for key in ['samsung','xiaomi','vivo','oppo','realme','infinix','tecno','appliances']:
   if not (ROOT/'raw'/f'{key}.html').exists():continue
   ss=soup(key)
   for box in ss.select('.productBox'):
    title=box.select_one('.p-title');price=box.select_one('.price-box');a=box.select_one('a[href]')
    if not title or not price or not a:continue
    name=title.get_text(' ',strip=True);url=a['href']
    if not url.startswith('https://priceoye.pk/'):continue
    v=num(price.get_text(' ',strip=True))
    add('retail-'+slug(url.split('/')[-1]),name,'electronics',v,'device; verify variant',TODAY,url,kind='Retailer listing · observed',note='Listing price; exact storage, warranty, stock and delivery must be checked at the retailer.')
    count+=1
  if not count:raise ValueError('No retail listing parsed')
 attempt('PriceOye published listings',retail)
 def weekly():
  file=ROOT/'assets/pbs-weekly.xlsx';w=load_workbook(file,data_only=True,read_only=True);rows=list(w['Appendix-A'].values);cities={};d=None
  for r in rows:
   if isinstance(r[3],str)and re.search(r'\(\d+\)',r[3]):cities={j:re.sub(r'\s*\(\d+\)','',r[j]).strip() for j in range(3,len(r)-2,3) if isinstance(r[j],str) and re.search(r'\(\d+\)',r[j])}
   for v in r:
    if isinstance(v,str) and 'PRICES ON' in v:
     m=re.search(r'\d{2}-\d{2}-\d{4}',v)
     if m:d=datetime.datetime.strptime(m[0],'%d-%m-%Y').date().isoformat()
   if not d or not isinstance(r[0],(int,float))or not isinstance(r[1],str)or not isinstance(r[2],str):continue
   if not 1<=r[0]<=51:continue
   name=r[1].strip();unit=r[2].strip()
   for j,city in cities.items():
    price=r[j+1]
    if isinstance(price,(int,float)) and price>0:
     add('pbs-weekly-'+str(int(r[0]))+'-'+slug(city),name,classified(name),price,unit,d,'https://www.pbs.gov.pk/price-statistics/',city=city,kind='PBS surveyed weekly average',period='Week ending '+d,min=r[j] if isinstance(r[j],(int,float))else None,max=r[j+2]if isinstance(r[j+2],(int,float))else None,note='Surveyed market average, not a government price ceiling. Date is the week ending; source workbook linked in the source library.')
 attempt('PBS weekly city prices',weekly)
 def construction():
  rows=list(load_workbook(ROOT/'assets/pbs-construction.xlsx',data_only=True,read_only=True).active.values)
  month=rows[1][0].replace(',','').strip();dt=datetime.datetime.strptime(month,'%B %Y');import calendar
  d=dt.replace(day=calendar.monthrange(dt.year,dt.month)[1]).date().isoformat()
  specs=[('Bricks · first class','1000 bricks'),('Cement blocks · 6×8×12 inches','100 blocks'),('Bajri / aggregate','100 CFT truck'),('Sand · black pit','100 CFT truck'),('Cement · local','50 kg bag'),('Iron bars · 1/2 inch, grade 40','tonne'),('Synthetic enamel · ICI','3.64 litre'),('Painter','day'),('Mason','day'),('Labour · unskilled','day'),('Plumber','day'),('Electrician','point')]
  for r in rows[3:]:
   if not isinstance(r[0],str):continue
   for j,(name,unit)in enumerate(specs,1):
    if len(r)<=j or not isinstance(r[j],(int,float))or r[j]<=0:continue
    add('pbs-construction-'+str(j)+'-'+slug(r[0]),name,'construction' if j<8 else 'services',r[j],unit,d,'https://www.pbs.gov.pk/price-statistics/',city=r[0].strip(),kind='PBS monthly surveyed average',period=month,note='Average for '+month+'; date marks month end. Not a daily dealer quote.')
 attempt('PBS construction and labour',construction)
 # Human-verified image transcriptions live separately so image changes never silently revise figures.
 p=ROOT/'data/manual-quotes.json'
 if p.exists():
  for r in json.loads(p.read_text()):add(**r)
 for id,r in records.items():
  pts={x['date']:x for x in hist.get(id,[])};pts[r['date']]={'date':r['date'],'price':r['price'],'source':r['source']};hist[id]=sorted(pts.values(),key=lambda x:x['date'])[-400:]
 (ROOT/'data/prices.json').write_text(json.dumps(list(records.values()),ensure_ascii=False,indent=2)+'\n')
 (ROOT/'data/history.json').write_text(json.dumps(hist,ensure_ascii=False,separators=(',',':'))+'\n')
 report={'fetches':json.loads((ROOT/'data/fetch-status.json').read_text()) if (ROOT/'data/fetch-status.json').exists() else [],'checked_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'sources':statuses,'quotes':len(records),'manual_queue':['Verify new image rates before transcribing; source sheets refresh automatically','Unsupported sources retain their dated quotes']}
 (ROOT/'data/update-report.json').write_text(json.dumps(report,indent=2)+'\n')
 print(json.dumps(report,indent=2))
if __name__=='__main__':ingest()
